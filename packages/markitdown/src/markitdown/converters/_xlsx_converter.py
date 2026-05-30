import sys
from typing import BinaryIO, Any
from ._html_converter import HtmlConverter
from .._base_converter import DocumentConverter, DocumentConverterResult
from .._exceptions import MissingDependencyException, MISSING_DEPENDENCY_MESSAGE
from .._stream_info import StreamInfo

# Try loading optional (but in this case, required) dependencies
# Save reporting of any exceptions for later
_xlsx_dependency_exc_info = None
try:
    import pandas as pd
    import openpyxl
except ImportError:
    _xlsx_dependency_exc_info = sys.exc_info()

_xls_dependency_exc_info = None
try:
    import pandas as pd  # noqa: F811
    import xlrd  # noqa: F401
except ImportError:
    _xls_dependency_exc_info = sys.exc_info()

ACCEPTED_XLSX_MIME_TYPE_PREFIXES = [
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
]
ACCEPTED_XLSX_FILE_EXTENSIONS = [".xlsx"]

ACCEPTED_XLS_MIME_TYPE_PREFIXES = [
    "application/vnd.ms-excel",
    "application/excel",
]
ACCEPTED_XLS_FILE_EXTENSIONS = [".xls"]


class XlsxConverter(DocumentConverter):
    """
    Converts XLSX files to Markdown, with each sheet presented as a separate Markdown table.
    """

    def __init__(self):
        super().__init__()
        self._html_converter = HtmlConverter()

    def accepts(
        self,
        file_stream: BinaryIO,
        stream_info: StreamInfo,
        **kwargs: Any,  # Options to pass to the converter
    ) -> bool:
        mimetype = (stream_info.mimetype or "").lower()
        extension = (stream_info.extension or "").lower()

        if extension in ACCEPTED_XLSX_FILE_EXTENSIONS:
            return True

        for prefix in ACCEPTED_XLSX_MIME_TYPE_PREFIXES:
            if mimetype.startswith(prefix):
                return True

        return False

    def convert(
        self,
        file_stream: BinaryIO,
        stream_info: StreamInfo,
        **kwargs: Any,  # Options to pass to the converter
    ) -> DocumentConverterResult:
        # Check the dependencies
        if _xlsx_dependency_exc_info is not None:
            raise MissingDependencyException(
                MISSING_DEPENDENCY_MESSAGE.format(
                    converter=type(self).__name__,
                    extension=".xlsx",
                    feature="xlsx",
                )
            ) from _xlsx_dependency_exc_info[
                1
            ].with_traceback(  # type: ignore[union-attr]
                _xlsx_dependency_exc_info[2]
            )

        sheets = pd.read_excel(file_stream, sheet_name=None, engine="openpyxl")
        self._apply_excel_number_formats(file_stream, sheets)

        md_content = ""
        for s in sheets:
            md_content += f"## {s}\n"
            html_content = sheets[s].to_html(index=False)
            md_content += (
                self._html_converter.convert_string(
                    html_content, **kwargs
                ).markdown.strip()
                + "\n\n"
            )

        return DocumentConverterResult(markdown=md_content.strip())

    def _apply_excel_number_formats(
        self, file_stream: BinaryIO, sheets: dict[str, Any]
    ) -> None:
        cur_pos = file_stream.tell()
        try:
            file_stream.seek(0)
            workbook = openpyxl.load_workbook(
                file_stream, data_only=True, read_only=True
            )
            for sheet_name, frame in sheets.items():
                if sheet_name not in workbook.sheetnames:
                    continue

                worksheet = workbook[sheet_name]
                frame = frame.astype(object)
                sheets[sheet_name] = frame
                for row_idx in range(len(frame.index)):
                    for col_idx in range(len(frame.columns)):
                        cell = worksheet.cell(row=row_idx + 2, column=col_idx + 1)
                        formatted = self._format_currency_cell(
                            frame.iat[row_idx, col_idx], cell.number_format
                        )
                        if formatted is not None:
                            frame.iat[row_idx, col_idx] = formatted
        finally:
            file_stream.seek(cur_pos)

    def _format_currency_cell(self, value: Any, number_format: str) -> str | None:
        if not isinstance(value, (int, float)):
            return None

        symbol = self._currency_symbol(number_format)
        if symbol is None:
            return None

        if ".00" in number_format or "0.00" in number_format:
            amount = f"{value:,.2f}"
        elif float(value).is_integer():
            amount = f"{int(value):,}"
        else:
            amount = f"{value:,}"

        return f"{symbol}{amount}"

    def _currency_symbol(self, number_format: str) -> str | None:
        for symbol in ["$", "€", "£", "¥", "₹"]:
            if symbol in number_format:
                return symbol
        return None


class XlsConverter(DocumentConverter):
    """
    Converts XLS files to Markdown, with each sheet presented as a separate Markdown table.
    """

    def __init__(self):
        super().__init__()
        self._html_converter = HtmlConverter()

    def accepts(
        self,
        file_stream: BinaryIO,
        stream_info: StreamInfo,
        **kwargs: Any,  # Options to pass to the converter
    ) -> bool:
        mimetype = (stream_info.mimetype or "").lower()
        extension = (stream_info.extension or "").lower()

        if extension in ACCEPTED_XLS_FILE_EXTENSIONS:
            return True

        for prefix in ACCEPTED_XLS_MIME_TYPE_PREFIXES:
            if mimetype.startswith(prefix):
                return True

        return False

    def convert(
        self,
        file_stream: BinaryIO,
        stream_info: StreamInfo,
        **kwargs: Any,  # Options to pass to the converter
    ) -> DocumentConverterResult:
        # Load the dependencies
        if _xls_dependency_exc_info is not None:
            raise MissingDependencyException(
                MISSING_DEPENDENCY_MESSAGE.format(
                    converter=type(self).__name__,
                    extension=".xls",
                    feature="xls",
                )
            ) from _xls_dependency_exc_info[
                1
            ].with_traceback(  # type: ignore[union-attr]
                _xls_dependency_exc_info[2]
            )

        sheets = pd.read_excel(file_stream, sheet_name=None, engine="xlrd")
        md_content = ""
        for s in sheets:
            md_content += f"## {s}\n"
            html_content = sheets[s].to_html(index=False)
            md_content += (
                self._html_converter.convert_string(
                    html_content, **kwargs
                ).markdown.strip()
                + "\n\n"
            )

        return DocumentConverterResult(markdown=md_content.strip())
